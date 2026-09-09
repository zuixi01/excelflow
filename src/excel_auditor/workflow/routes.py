from __future__ import annotations

import os
import re
import threading
from functools import lru_cache
from pathlib import Path
from urllib.parse import urlsplit

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, Header, Request, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import Field
from sqlalchemy import select
from starlette.background import BackgroundTask

from ..models import StrictModel
from .ai import AIClient, AISettings
from .excel import MAX_UPLOAD
from .models import EditRequest, MappingRequest, RevisionRequest, TemplateEdit, WorkflowError
from .repository import WorkflowStore, objects
from .tasks import TaskService
from .templates import TemplateService


class NewTemplate(StrictModel):
    source_id: str | None = None
    name: str = Field(default="新业务模板", min_length=1, max_length=100)


class DeleteTemplates(StrictModel):
    ids: list[str] = Field(min_length=1, max_length=100)


class DeleteTasks(StrictModel):
    ids: list[str] = Field(min_length=1, max_length=100)


class LearnRequest(RevisionRequest):
    instructions: str = Field(default="", max_length=6000)


class ReturnRequest(StrictModel):
    preview_id: str
    delete_missing: bool = False


class ExportRequest(RevisionRequest):
    output_template_id: str | None = Field(default=None, min_length=1, max_length=64)


def guard_origin(request: Request):
    if request.method not in {"GET", "HEAD", "OPTIONS"} and (origin := request.headers.get("origin")):
        origin_host = urlsplit(origin).hostname
        expected = request.url.hostname
        if origin_host != expected and not {origin_host, expected} <= {"localhost", "127.0.0.1", "::1"}:
            raise WorkflowError("请求来源不匹配，请从本项目页面操作。", 403)


def require_admin(request: Request):
    allowed = {name.strip() for name in os.environ.get("EXCEL_AUDITOR_TEMPLATE_ADMINS", "local,service").split(",") if name.strip()}
    if request.state.user_id not in allowed:
        raise WorkflowError("当前账户没有模板发布或模型配置权限。", 403)


async def read_upload(file: UploadFile) -> bytes:
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise WorkflowError("当前流程支持 .xlsx 文件。", 415)
    content = bytearray()
    while chunk := await file.read(1024 * 1024):
        content.extend(chunk)
        if len(content) > MAX_UPLOAD:
            raise WorkflowError("上传文件不能超过 20 MiB。", 413)
    return bytes(content)


def build_router(data_root: Path, database=None, task_queue=None) -> APIRouter:
    router = APIRouter(prefix="/api/v1", dependencies=[Depends(guard_origin)])

    instance = None
    initialization_lock = threading.Lock()

    def store():
        nonlocal instance
        with initialization_lock:
            if instance is None:
                instance = WorkflowStore(data_root / "workflow", database.engine if database else None)
        return instance

    def tasks():
        return TaskService(store())

    def templates():
        return TemplateService(store())

    def dispatch(background: BackgroundTasks, operation: str, object_id: str, tenant: str):
        if task_queue:
            try:
                task_queue.queue.enqueue("excel_auditor.workflow.routes.run_background", str(data_root / "workflow"),
                    operation, object_id, tenant, job_timeout=300, result_ttl=3600, failure_ttl=86400)
            except Exception as exc:
                tasks().fail(object_id, tenant, WorkflowError("任务队列暂不可用，请稍后重试。", 503),
                             "task" if operation == "parse_task" else "export" if operation == "export" else "template")
                raise WorkflowError("任务队列暂不可用，请稍后重试。", 503) from exc
        else:
            function = {"parse_task": tasks().parse, "export": tasks().generate,
                        "parse_template": templates().parse, "learn": templates().learn}[operation]
            background.add_task(function, object_id, tenant)

    @router.get("/workflow-definitions")
    def definitions(request: Request):
        return templates().definitions(request.state.tenant_id)

    @router.get("/workflow-examples/source")
    def example():
        from io import BytesIO
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "商家商品"
        sheet.append(["商品名称", "货号", "商品分类", "商品品牌", "销售价", "库存", "销售单位"])
        sheet.append(["示例台灯", "0001", "照明", "示例品牌", "129.50", "20", "盏"])
        sheet.append(["示例吊灯", "0002", "照明", "示例品牌", "299.00", "12", "盏"])
        output = BytesIO()
        book.save(output)
        book.close()
        return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @router.get("/workflow-tasks")
    def list_tasks(request: Request, search: str = "", status: str = "", offset: int = 0, limit: int = 30):
        return tasks().list(request.state.tenant_id, search[:200], status, max(0, offset), max(1, min(100, limit)))

    @router.delete("/workflow-tasks")
    def delete_tasks(request: Request, body: DeleteTasks):
        return tasks().delete(request.state.tenant_id, body.ids)

    @router.post("/workflow-tasks", status_code=202)
    async def create_task(request: Request, background: BackgroundTasks, excel_file: UploadFile = File(...),
                          template_id: str = Form(...), name: str = Form(""), idempotency_key: str | None = Header(None, alias="Idempotency-Key")):
        content = await read_upload(excel_file)
        filename = re.split(r"[/\\]", excel_file.filename or "原表.xlsx")[-1][:150]
        obj, created = tasks().create(request.state.tenant_id, request.state.user_id, template_id, name[:100], filename, content, idempotency_key)
        if created:
            dispatch(background, "parse_task", obj["id"], request.state.tenant_id)
        return obj

    @router.get("/workflow-tasks/{task_id}")
    def task(request: Request, task_id: str):
        return tasks().get(task_id, request.state.tenant_id)

    @router.post("/workflow-tasks/{task_id}/mapping")
    def mapping(request: Request, task_id: str, body: MappingRequest):
        return tasks().confirm_mapping(task_id, request.state.tenant_id, request.state.user_id, body)

    @router.post("/workflow-tasks/{task_id}/mapping/reopen")
    def reopen_mapping(request: Request, task_id: str, body: RevisionRequest):
        return tasks().reopen_mapping(task_id, request.state.tenant_id, request.state.user_id, body.base_revision)

    @router.get("/workflow-tasks/{task_id}/maintenance")
    def maintenance(request: Request, task_id: str, offset: int = 0, limit: int = 30, search: str = "", issues_only: bool = False, sort: str = "", descending: bool = False):
        return tasks().rows(task_id, request.state.tenant_id, max(0, offset), max(1, min(100, limit)), search[:200], issues_only, sort, descending)

    @router.patch("/workflow-tasks/{task_id}/maintenance")
    def edit(request: Request, task_id: str, body: EditRequest):
        return tasks().edit(task_id, request.state.tenant_id, request.state.user_id, body)

    @router.post("/workflow-tasks/{task_id}/maintenance/validate")
    def validate(request: Request, task_id: str, body: RevisionRequest):
        return tasks().edit(task_id, request.state.tenant_id, request.state.user_id, EditRequest(base_revision=body.base_revision))

    @router.get("/workflow-tasks/{task_id}/maintenance/download")
    def download_maintenance(request: Request, task_id: str):
        path = tasks().maintenance_file(task_id, request.state.tenant_id)
        return FileResponse(path, filename="商品维护表.xlsx", background=BackgroundTask(path.unlink, missing_ok=True))

    @router.post("/workflow-tasks/{task_id}/maintenance/returns")
    async def preview_return(request: Request, task_id: str, excel_file: UploadFile = File(...)):
        return tasks().preview_return(task_id, request.state.tenant_id, await read_upload(excel_file))

    @router.post("/workflow-tasks/{task_id}/maintenance/returns/apply")
    def apply_return(request: Request, task_id: str, body: ReturnRequest):
        return tasks().apply_return(task_id, request.state.tenant_id, request.state.user_id, body.preview_id, body.delete_missing)

    @router.post("/workflow-tasks/{task_id}/exports", status_code=202)
    def create_export(request: Request, background: BackgroundTasks, task_id: str, body: ExportRequest,
                      idempotency_key: str | None = Header(None, alias="Idempotency-Key")):
        obj, created = tasks().create_export(task_id, request.state.tenant_id, body.base_revision,
                                             idempotency=idempotency_key, output_template_id=body.output_template_id)
        if created:
            dispatch(background, "export", obj["id"], request.state.tenant_id)
        return obj

    @router.get("/workflow-tasks/{task_id}/exports/{run_id}/files/{filename}")
    def artifact(request: Request, task_id: str, run_id: str, filename: str):
        return FileResponse(tasks().export_file(task_id, run_id, request.state.tenant_id, filename), filename=filename)

    @router.get("/workflow-tasks/{task_id}/exports/{run_id}/preview/{filename}")
    def preview_output(request: Request, task_id: str, run_id: str, filename: str):
        from openpyxl import load_workbook
        from .excel import cell_text
        path = tasks().export_file(task_id, run_id, request.state.tenant_id, filename)
        if path.suffix != ".xlsx":
            raise WorkflowError("仅支持预览 Excel 文件。")
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return {"name": filename, "sheets": [{"name": sheet.title,
                "rows": [[cell_text(cell) for cell in row] for row in sheet.iter_rows(max_row=min(12, sheet.max_row), values_only=True)]}
                for sheet in book.worksheets if sheet.sheet_state == "visible"]}
        finally:
            book.close()

    @router.post("/workflow-tasks/{task_id}/cancel")
    def cancel(request: Request, task_id: str, body: RevisionRequest):
        return tasks().cancel(task_id, request.state.tenant_id, body.base_revision)

    @router.post("/workflow-tasks/{task_id}/retry", status_code=202)
    def retry(request: Request, background: BackgroundTasks, task_id: str, body: RevisionRequest):
        with store().engine.begin() as con:
            obj = store().claim(con, task_id, request.state.tenant_id, body.base_revision)
            if obj["status"] not in {"failed", "cancelled"} or obj["payload"].get("record_count"):
                raise WorkflowError("当前任务无需重新解析。", 409)
            obj["status"] = "processing"
            obj["payload"].pop("error", None)
            store().save(con, obj)
        dispatch(background, "parse_task", task_id, request.state.tenant_id)
        return store().public(obj)

    @router.get("/workflow-templates")
    def list_templates(request: Request):
        return templates().list(request.state.tenant_id)

    @router.delete("/workflow-templates", dependencies=[Depends(require_admin)])
    def delete_templates(request: Request, body: DeleteTemplates):
        return templates().delete(request.state.tenant_id, body.ids)

    @router.post("/workflow-templates", dependencies=[Depends(require_admin)], status_code=201)
    def create_template(request: Request, body: NewTemplate):
        return templates().create(request.state.tenant_id, body.source_id, body.name)

    @router.post("/workflow-templates/import", dependencies=[Depends(require_admin)], status_code=202)
    async def import_template(request: Request, background: BackgroundTasks, excel_file: UploadFile = File(...), name: str = Form("新商品模板"), purpose: str = Form("output")):
        obj = templates().upload(request.state.tenant_id, name[:100], await read_upload(excel_file), purpose)
        dispatch(background, "parse_template", obj["id"], request.state.tenant_id)
        return obj

    @router.post("/workflow-templates/{template_id}/sources", dependencies=[Depends(require_admin)], status_code=202)
    async def attach_template_source(request: Request, background: BackgroundTasks, template_id: str, excel_file: UploadFile = File(...), purpose: str = Form(...)):
        obj = templates().attach_source(template_id, request.state.tenant_id, await read_upload(excel_file), purpose)
        dispatch(background, "parse_template", obj["id"], request.state.tenant_id)
        return obj

    @router.post("/workflow-templates/{template_id}/retry", dependencies=[Depends(require_admin)], status_code=202)
    def retry_template(request: Request, background: BackgroundTasks, template_id: str, body: RevisionRequest):
        obj = templates().retry(template_id, request.state.tenant_id, body.base_revision)
        dispatch(background, "parse_template", obj["id"], request.state.tenant_id)
        return obj

    @router.post("/workflow-templates/{template_id}/reparse", dependencies=[Depends(require_admin)], status_code=202)
    def reparse_template(request: Request, background: BackgroundTasks, template_id: str, body: RevisionRequest):
        obj = templates().reparse(template_id, request.state.tenant_id, body.base_revision)
        dispatch(background, "parse_template", obj["id"], request.state.tenant_id)
        return obj

    @router.get("/workflow-templates/{template_id}")
    def template(request: Request, template_id: str):
        return templates().get(template_id, request.state.tenant_id)

    @router.put("/workflow-templates/{template_id}", dependencies=[Depends(require_admin)])
    def update_template(request: Request, template_id: str, body: TemplateEdit):
        return templates().update(template_id, request.state.tenant_id, body)

    @router.post("/workflow-templates/{template_id}/learn", dependencies=[Depends(require_admin)], status_code=202)
    def learn(request: Request, background: BackgroundTasks, template_id: str, body: LearnRequest):
        obj = templates().ai_suggest(template_id, request.state.tenant_id, body.base_revision, body.instructions)
        dispatch(background, "learn", template_id, request.state.tenant_id)
        return obj

    @router.post("/workflow-templates/{template_id}/trial", dependencies=[Depends(require_admin)])
    async def trial(request: Request, template_id: str, base_revision: int = Form(...), sample_file: UploadFile | None = File(None)):
        content = await read_upload(sample_file) if sample_file else None
        return templates().trial(template_id, request.state.tenant_id, base_revision, content)

    @router.get("/workflow-templates/{template_id}/trial/{index}")
    def trial_file(request: Request, template_id: str, index: int):
        obj = templates().get(template_id, request.state.tenant_id)
        artifacts = obj["payload"].get("test_result", {}).get("artifacts", [])
        if index < 0 or index >= len(artifacts):
            raise WorkflowError("试运行文件不存在。", 404)
        return FileResponse(store().directory(template_id) / f"trial-{index}.xlsx", filename=artifacts[index]["name"])

    @router.post("/workflow-templates/{template_id}/publish", dependencies=[Depends(require_admin)])
    def publish(request: Request, template_id: str, body: RevisionRequest):
        return templates().publish(template_id, request.state.tenant_id, body.base_revision)

    @router.get("/workflow-settings")
    def settings(request: Request):
        allowed = {name.strip() for name in os.environ.get("EXCEL_AUDITOR_TEMPLATE_ADMINS", "local,service").split(",")}
        return {"ai": AIClient(store().root, request.state.tenant_id).settings(),
                "can_manage_templates": request.state.user_id in allowed,
                "limits": {"upload_mib": 20, "rows": 20000, "columns": 200},
                "storage": "database" if database else "local"}

    @router.put("/workflow-settings/ai", dependencies=[Depends(require_admin)])
    def save_ai(request: Request, body: AISettings):
        return AIClient(store().root, request.state.tenant_id).save(body)

    @router.post("/workflow-settings/ai/models", dependencies=[Depends(require_admin)])
    def models(request: Request):
        return AIClient(store().root, request.state.tenant_id).models()

    return router


def run_background(root: str, operation: str, object_id: str, tenant: str):
    from sqlalchemy import create_engine
    engine = create_engine(os.environ["DATABASE_URL"], pool_pre_ping=True) if os.environ.get("DATABASE_URL") else None
    store = WorkflowStore(Path(root), engine)
    try:
        function = {"parse_task": TaskService(store).parse, "export": TaskService(store).generate,
                    "parse_template": TemplateService(store).parse, "learn": TemplateService(store).learn}.get(operation)
        if function is None:
            raise WorkflowError("不支持的后台任务。")
        function(object_id, tenant)
    finally:
        store.engine.dispose()
