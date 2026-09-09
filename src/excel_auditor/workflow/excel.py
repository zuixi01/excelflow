from __future__ import annotations

import hashlib
import itertools
import json
import re
import unicodedata
import zipfile
from copy import copy
from datetime import date, datetime
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..engine import _validate
from ..models import ColumnRule
from ..normalization import excel_numeric_write_safe, normalized_uniqueness_key, parse_value
from ..product_workflow.fixed_template import match_fixed_field, normalize_header
from ..product_workflow.sheet_discovery import _header_depth, flatten_headers, merged_value_map
from ..workbook import _validate_xml_part, _inspect_sheet_xml_structure
from .models import MaintenanceField, OutputFile, WorkflowConfig, WorkflowError


MAX_UPLOAD = 20 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 200
MAX_CELLS = 2_000_000


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, default=str, separators=(",", ":"))
    return str(value)


def safe_book(path: Path):
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > MAX_UPLOAD:
        raise WorkflowError("SOP 支持 20 MiB 以内的 .xlsx；请先将宏工作簿另存为普通数据表。")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 80 * 1024 * 1024:
                raise WorkflowError("解压后工作簿过大，超出当前 SOP 容量。")
            names = [item.filename.lower() for item in entries]
            if len(set(names)) != len(names) or any(".." in name.split("/") or name.startswith(("/", "\\")) for name in names):
                raise WorkflowError("工作簿包含无效文件条目。")
            unsupported = ("vbaproject", "xl/activex", "xl/embeddings", "xl/externallinks", "xl/charts", "xl/media", "xl/pivottables", "xl/connections.xml")
            if any(any(feature in name for feature in unsupported) for name in names):
                raise WorkflowError("文件包含图片、宏、控件、透视表或外部连接；当前 SOP 不会静默丢弃，请提供普通数据模板。")
            total_cells = 0
            for item in entries:
                if item.flag_bits & 1:
                    raise WorkflowError("请先解除工作簿文件加密。")
                if item.filename.lower().endswith((".xml", ".rels")):
                    _validate_xml_part(archive, item)
                if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                    _, bounds = _inspect_sheet_xml_structure(archive, item)
                    if bounds[0] > MAX_ROWS + 100 or bounds[1] > MAX_COLUMNS:
                        raise WorkflowError("当前 SOP 每表支持 20,000 条记录、200 列。")
                    total_cells += bounds[0] * bounds[1]
            if total_cells > MAX_CELLS:
                raise WorkflowError("工作簿数据区域超过 200 万单元格。")
        book = load_workbook(path, keep_links=False)
        dimensions = [_content_bounds(sheet) for sheet in book.worksheets]
        if len(book.worksheets) > 50 or any(row > MAX_ROWS + 100 or column > MAX_COLUMNS for row, column in dimensions) or sum(row * column for row, column in dimensions) > MAX_CELLS:
            book.close()
            raise WorkflowError("工作表数量或实际数据区域超出当前容量，请拆分后重试。")
        return book
    except WorkflowError:
        raise
    except Exception as exc:
        raise WorkflowError("无法解析工作簿，请检查文件是否完整且为有效 .xlsx。") from exc


def _content_bounds(sheet) -> tuple[int, int]:
    """Return the real data extent, ignoring Excel's trailing empty formatting."""
    content = [cell for cell in sheet._cells.values() if cell.value is not None]
    if not content:
        return 0, 0
    return max(cell.row for cell in content), max(cell.column for cell in content)


def field_matches(header: str, fields: list[MaintenanceField]) -> list[str]:
    token = normalize_header(header)
    matches = [f.rule.name for f in fields if token and token in {
        normalize_header(f.rule.name), normalize_header(f.rule.title), *(normalize_header(a) for a in f.rule.aliases)
    }]
    fixed = match_fixed_field(header)
    if not matches and fixed in {f.rule.name for f in fields}:
        matches = [fixed]
    return matches


def discover(path: Path, fields: list[MaintenanceField], *, template: bool = False) -> list[dict]:
    book = safe_book(path)
    result = []
    count = 0
    try:
        for sheet in book.worksheets:
            if sheet.sheet_state != "visible":
                continue
            last_row, last_column = _content_bounds(sheet)
            if not last_row or not last_column:
                continue
            rows = []
            for row in sheet.iter_rows(max_row=last_row, max_col=last_column):
                values = []
                for cell in row:
                    if cell.data_type == "f" and not template:
                        raise WorkflowError(f"{sheet.title}!{cell.coordinate} 含公式，请先在原表中粘贴为值再导入。")
                    value = cell_text(cell.value)
                    if cell.data_type == "n" and isinstance(cell.value, int) and re.fullmatch(r"0{2,20}", cell.number_format or ""):
                        value = str(cell.value).zfill(len(cell.number_format))
                    values.append(value)
                rows.append(values)
            while rows and not any(rows[-1]):
                rows.pop()
            if not rows:
                continue
            # Prefer a row that starts a merged, multi-level header when present.  A
            # maintenance template may use entirely custom field names, so known
            # product-field aliases alone cannot decide where its header starts.
            scores = []
            for index, row in enumerate(rows[:50]):
                known = sum(bool(field_matches(value, fields)) for value in row if value)
                labels = sum(bool(value) and not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", value) for value in row)
                structural = 0
                if template and index + 1 < len(rows):
                    for region in sheet.merged_cells.ranges:
                        if region.min_row != index + 1 or region.max_col <= region.min_col:
                            continue
                        child_count = sum(bool(rows[index + 1][column - 1]) for column in range(region.min_col, region.max_col + 1)
                                          if column <= len(rows[index + 1]))
                        if child_count >= 2:
                            structural += 1
                scores.append((structural * 10000 + known * 100 + min(labels, 20), -index))
            header_row = -max(scores)[1] + 1
            depth = _header_depth(sheet, header_row, len(rows), last_column)
            merged = merged_value_map(sheet, last_row=len(rows), last_column=last_column)
            headers = flatten_headers(sheet, merged, header_row=header_row, header_depth=depth, last_column=last_column)
            for (r, c), value in merged.items():
                if r <= len(rows) and c <= len(rows[r - 1]):
                    rows[r - 1][c - 1] = cell_text(value)
            columns = []
            for index, header in enumerate(headers, start=1):
                matches = field_matches(header, fields)
                columns.append({"column": index, "header": header or f"未命名列 {index}", "candidates": matches,
                                "field": matches[0] if len(matches) == 1 else "",
                                "example": next((r[index - 1] for r in rows[header_row + depth - 1:] if index <= len(r) and r[index - 1]), "")[:200]})
            count += max(0, len(rows) - header_row - depth + 1)
            if count > MAX_ROWS:
                raise WorkflowError("一批数据合计最多 20,000 条记录，请拆分后导入。")
            result.append({"sheet": sheet.title, "header_row": header_row, "header_depth": depth,
                           "row_count": max(0, len(rows) - header_row - depth + 1), "columns": columns,
                           "preview": rows[:min(len(rows), 8)], "rows": rows,
                           "widths": [sheet.column_dimensions[get_column_letter(i)].width or 18 for i in range(1, last_column + 1)]})
    finally:
        book.close()
    if not result:
        raise WorkflowError("未找到可用的可见工作表。")
    return result


def transform(values: list[str], operation: str, argument: str = "", part: int = 0) -> str:
    value = values[0] if values else ""
    if operation == "constant":
        return argument
    if operation == "concat":
        return argument.join(v for v in values if v)
    if operation == "trim":
        return value.strip()
    if operation == "split":
        parts = value.split(argument)
        return parts[part] if part < len(parts) else ""
    if operation == "multiply" and value.strip():
        try:
            with localcontext() as ctx:
                ctx.prec = max(50, len(value) + len(argument) + 10)
                number = Decimal(value) * Decimal(argument)
                if not number.is_finite() or abs(number.adjusted()) > 1000:
                    raise ValueError
                return str(number)
        except Exception as exc:
            raise WorkflowError("数值换算失败，请检查原值和换算系数。") from exc
    return value


def validate_records(rows: list[dict], config: WorkflowConfig) -> tuple[int, int]:
    uniqueness: dict[tuple, list[dict]] = {}
    errors = 0
    warnings = 0
    for row in rows:
        row["issues"] = []
        parsed = {}
        for field in config.fields:
            rule = field.rule
            value = row["values_json"].get(rule.name, "")
            if value is None:
                value = ""
            item = parse_value(value, rule)
            parsed[rule.name] = item
            type_labels = {"decimal": "有效数值，例如 129.50", "integer": "整数", "date": "有效日期，例如 2026-09-08",
                           "datetime": "有效日期时间", "boolean": "是或否", "enum": "模板允许的选项", "json": "有效 JSON"}
            error = (f"请填写{type_labels.get(rule.type.value, '符合字段类型的值')}。") if not item.valid else _validate(item, rule)
            if isinstance(value, str) and value.startswith("="):
                error = "维护值不能是公式，请填写实际值；公式请在输出模板中配置。"
            if error:
                row["issues"].append({"field": rule.name, "title": rule.title, "message": error, "severity": "error"})
            elif item.valid:
                row["values_json"][rule.name] = cell_text(item.normalized)
                if rule.validation.unique and item.normalized is not None:
                    key = normalized_uniqueness_key(item, rule)
                    uniqueness.setdefault((rule.name, key), []).append(row)
        if config.unique_key:
            keys = [parsed[name] for name in config.unique_key]
            if all(item.valid and item.normalized is not None for item in keys):
                field_by_id = {f.rule.name: f.rule for f in config.fields}
                key = tuple(normalized_uniqueness_key(parsed[name], field_by_id[name]) for name in config.unique_key)
                uniqueness.setdefault(("business_key", key), []).append(row)
            else:
                row["issues"].append({"field": config.unique_key[0], "message": "业务唯一键不能为空且必须有效", "severity": "error"})
        for check in config.cross_checks:
            left, right = parsed[check.left], parsed[check.right]
            failed = False
            if check.kind == "less_equal" and left.valid and right.valid and left.normalized is not None and right.normalized is not None:
                failed = left.normalized > right.normalized
            elif check.kind == "conditional_required":
                failed = left.valid and cell_text(left.normalized) == check.equals and right.normalized is None
            if failed:
                row["issues"].append({"field": check.right, "message": check.message, "severity": "error"})
        row["search_text"] = " ".join(cell_text(v) for v in row["values_json"].values()).casefold()
    for (field, _), matches in uniqueness.items():
        if len(matches) > 1:
            for row in matches:
                row["issues"].append({"field": config.unique_key[0] if field == "business_key" else field,
                                      "message": "存在重复的业务键或唯一字段值", "severity": "error"})
    for row in rows:
        errors += sum(issue["severity"] == "error" for issue in row["issues"])
        warnings += sum(issue["severity"] == "warning" for issue in row["issues"])
    return errors, warnings


def literal(cell, value: Any):
    value = cell_text(value)
    if len(value) > 32767:
        raise WorkflowError("单元格文本超过 Excel 32,767 字符限制。")
    cell.value = value or None
    if value:
        cell.data_type = "s"


def typed_cell(cell, value: str, rule: ColumnRule | None):
    if rule is None or rule.type.value in {"string", "phone", "id_code", "postal_code", "enum", "fuzzy_string", "set", "json"}:
        literal(cell, value)
        return
    parsed = parse_value(value, rule)
    if not parsed.valid:
        raise WorkflowError(f"{rule.title} 输出类型不正确。")
    value = parsed.normalized
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        if not excel_numeric_write_safe(value):
            raise WorkflowError(f"{rule.title} 超过 Excel 安全数值精度，请改用文本类型或调整业务值。")
        cell.value = int(value) if value == int(value) else float(value)
    elif isinstance(value, datetime):
        if value.tzinfo is not None:
            # Do not erase timezone semantics; expose ISO text for zoned values.
            literal(cell, value.isoformat())
        else:
            cell.value = value
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = "yyyy-mm-dd"
    else:
        cell.value = value


_GRID_SIDE = Side(style="thin", color="D9D9D9")
_GRID_BORDER = Border(left=_GRID_SIDE, right=_GRID_SIDE, top=_GRID_SIDE, bottom=_GRID_SIDE)


def style_header(sheet, row: int, titles: list[str]):
    for column, title in enumerate(titles, start=1):
        cell = sheet.cell(row, column)
        literal(cell, title)
        style_header_cell(cell)
        sheet.column_dimensions[get_column_letter(column)].width = 16
    sheet.row_dimensions[row].height = 26
    sheet.freeze_panes = f"A{row + 1}"


def style_header_cell(cell):
    cell.fill = PatternFill(fill_type=None)
    cell.font = Font(name="Microsoft YaHei", bold=True, color="000000")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _GRID_BORDER


def style_data_cell(cell):
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = _GRID_BORDER


def display_width(value: Any) -> int:
    """Approximate Excel column width, accounting for full-width characters."""
    lines = cell_text(value).splitlines() or [""]
    return max(sum(2 if unicodedata.east_asian_width(char) in {"W", "F", "A"} else 1 for char in line) for line in lines)


def _header_paths(titles: list[str]) -> list[list[str]]:
    return [[part.strip() for part in title.split("/") if part.strip()] or [title] for title in titles]


def write_hierarchical_headers(sheet, titles: list[str], *, start_column: int = 1, start_row: int = 1) -> int:
    """Write slash-delimited field titles as a compact merged Excel header."""
    paths = _header_paths(titles)
    depth = max(map(len, paths), default=1)
    for offset, path in enumerate(paths):
        column = start_column + offset
        sheet.column_dimensions[get_column_letter(column)].width = max(10, min(24, max(map(len, path)) * 2 + 4))
        for level, title in enumerate(path, start=1):
            literal(sheet.cell(start_row + level - 1, column), title)
    for offset, path in enumerate(paths):
        column = start_column + offset
        for level, title in enumerate(path, start=1):
            if offset and len(paths[offset - 1]) > level - 1 and paths[offset - 1][:level] == path[:level]:
                continue
            end_column = column
            while end_column - start_column + 1 < len(paths):
                candidate = paths[end_column - start_column + 1]
                if len(candidate) <= level - 1 or candidate[:level] != path[:level]:
                    break
                end_column += 1
            end_row = start_row + depth - 1 if level == len(path) else start_row + level - 1
            current_row = start_row + level - 1
            for row in range(current_row, end_row + 1):
                for current_column in range(column, end_column + 1):
                    style_header_cell(sheet.cell(row, current_column))
            if end_row > current_row or end_column > column:
                sheet.merge_cells(start_row=current_row, start_column=column, end_row=end_row, end_column=end_column)
    for row in range(start_row, start_row + depth):
        sheet.row_dimensions[row].height = 26
    return depth


def write_maintenance(path: Path, task: dict, rows: list[dict]):
    config = WorkflowConfig.model_validate(task["payload"]["config"])
    book = Workbook()
    sheet = book.active
    sheet.title = "维护表"
    ids = [f.rule.name for f in config.fields]
    titles = [f.rule.title for f in config.fields]
    header_depth = write_hierarchical_headers(sheet, titles, start_column=2)
    literal(sheet.cell(1, 1), "_workflow_record_id")
    style_header_cell(sheet.cell(1, 1))
    if header_depth > 1:
        for row in range(2, header_depth + 1):
            style_header_cell(sheet.cell(row, 1))
        sheet.merge_cells(start_row=1, start_column=1, end_row=header_depth, end_column=1)
    sheet.column_dimensions["A"].hidden = True
    data_start_row = header_depth + 1
    for index, record in enumerate(rows, start=data_start_row):
        literal(sheet.cell(index, 1), record["id"])
        for column, field in enumerate(config.fields, start=2):
            # Maintenance transport deliberately uses text to preserve exact round trips.
            literal(sheet.cell(index, column), record["values_json"].get(field.rule.name, ""))
            style_data_cell(sheet.cell(index, column))
            if any(issue["field"] == field.rule.name for issue in record["issues"]):
                sheet.cell(index, column).fill = PatternFill("solid", fgColor="FCE8E6")
    sheet.freeze_panes = f"B{data_start_row}"
    sheet.auto_filter.ref = f"A{header_depth}:{get_column_letter(len(ids) + 1)}{max(header_depth, data_start_row + len(rows) - 1)}"
    meta = book.create_sheet("_workflow_meta")
    meta.sheet_state = "veryHidden"
    meta.append(["task_id", task["id"]])
    meta.append(["revision", task["revision"]])
    meta.append(["template_hash", task["payload"]["template_hash"]])
    meta.append(["fields", json.dumps(ids)])
    meta.append(["titles", json.dumps(titles, ensure_ascii=False)])
    meta.append(["header_depth", header_depth])
    book.save(path)
    book.close()


def read_maintenance(path: Path, task: dict) -> list[dict]:
    book = safe_book(path)
    try:
        if "_workflow_meta" not in book.sheetnames or "维护表" not in book.sheetnames:
            raise WorkflowError("维护文件缺少记录标识，请下载本任务维护表后编辑回传。")
        meta = {row[0]: row[1] for row in book["_workflow_meta"].iter_rows(values_only=True) if row[0]}
        if meta.get("task_id") != task["id"] or meta.get("template_hash") != task["payload"]["template_hash"]:
            raise WorkflowError("维护文件不属于当前任务或模板版本。", 409)
        if meta.get("revision") != task["revision"]:
            raise WorkflowError("维护文件已过期，请下载最新版本并重新应用修改。", 409, "REVISION_CONFLICT")
        config = WorkflowConfig.model_validate(task["payload"]["config"])
        ids = [f.rule.name for f in config.fields]
        if meta.get("fields") != json.dumps(ids):
            raise WorkflowError("维护字段元数据已被修改。")
        sheet = book["维护表"]
        header_depth = int(meta.get("header_depth", 1))
        titles = ["_workflow_record_id", *(f.rule.title for f in config.fields)]
        headers = flatten_headers(sheet, merged_value_map(sheet, last_row=header_depth, last_column=len(titles)),
                                  header_row=1, header_depth=header_depth, last_column=len(titles))
        if [normalize_header(header) for header in headers] != [normalize_header(title) for title in titles]:
            raise WorkflowError("维护表列结构已改变，请保留原列和表头。")
        rows = []
        for index, cells in enumerate(sheet.iter_rows(min_row=header_depth + 1), start=header_depth + 1):
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise WorkflowError(f"维护表第 {index} 行包含公式，请粘贴为值。")
            rows.append({"id": cell_text(cells[0].value), "values": {name: cell_text(cells[i + 1].value) for i, name in enumerate(ids)}})
        if len(rows) > MAX_ROWS:
            raise WorkflowError("维护记录超过 20,000 条。")
        return rows
    finally:
        book.close()


def write_output(path: Path, spec: OutputFile, config: WorkflowConfig, rows: list[dict], reference: Path | None = None) -> dict:
    book = safe_book(reference) if reference and reference.exists() else Workbook()
    if not reference and book.active:
        book.remove(book.active)
    rules = {field.rule.name: field.rule for field in config.fields}
    desired_names = {sheet.name for sheet in spec.sheets}
    # Preserve hidden validation lookup sheets but never leak sample business sheets.
    for existing in list(book.worksheets):
        if existing.sheet_state == "visible" and existing.title not in desired_names:
            book.remove(existing)
    stats = []
    seen_names = set()
    for plan in spec.sheets:
        grouped: dict[str, list[dict]] = {}
        for row in rows:
            key = row["values_json"].get(plan.group_by, "未分类") or "未分类" if plan.group_by else ""
            grouped.setdefault(key, []).append(row)
        if not grouped:
            grouped[""] = []
        for group, records_group in grouped.items():
            name = plan.name if not group else f"{plan.name}-{group}"
            name = re.sub(r"[\\/*?:\[\]]", "_", name)[:31]
            if name.casefold() in seen_names or name.startswith("_"):
                raise WorkflowError("分组后的工作表名重复，请缩短类目名称或调整输出方案。")
            seen_names.add(name.casefold())
            if len(seen_names) > 50:
                raise WorkflowError("输出拆分超过 50 个工作表。")
            if name in book.sheetnames:
                sheet = book[name]
            elif plan.name in book.sheetnames:
                sheet = book.copy_worksheet(book[plan.name])
                sheet.title = name
            else:
                sheet = book.create_sheet(name)
            if sheet.tables:
                raise WorkflowError("当前输出模板含 Excel Table，请转换为普通区域后学习模板。")
            if any(region.max_row >= plan.data_start_row for region in sheet.merged_cells.ranges):
                raise WorkflowError("输出模板数据区域存在合并单元格，请仅在表头保留合并区域。")
            styles = [copy(sheet.cell(plan.data_start_row, i)._style) for i in range(1, len(plan.columns) + 1)]
            headers = flatten_headers(sheet, merged_value_map(sheet, last_row=plan.data_start_row - 1, last_column=sheet.max_column), header_row=plan.header_row,
                                      header_depth=plan.data_start_row - plan.header_row, last_column=sheet.max_column)
            titles = [column.title for column in plan.columns]
            desired_depth = max(map(len, _header_paths(titles)), default=1)
            keep_headers = ([normalize_header(header) for header in headers] == [normalize_header(title) for title in titles]
                            and plan.data_start_row - plan.header_row >= desired_depth)
            header_depth = max(plan.data_start_row - plan.header_row, desired_depth)
            data_start_row = plan.header_row + header_depth
            if sheet.max_row >= plan.data_start_row:
                sheet.delete_rows(plan.data_start_row, sheet.max_row - plan.data_start_row + 1)
            if not keep_headers:
                for region in list(sheet.merged_cells.ranges):
                    if region.max_row >= plan.header_row:
                        sheet.unmerge_cells(str(region))
                for header_cells in sheet.iter_rows(min_row=plan.header_row, max_row=data_start_row - 1):
                    for cell in header_cells:
                        cell.value = None
                if sheet.max_column > len(plan.columns):
                    sheet.delete_cols(len(plan.columns) + 1, sheet.max_column - len(plan.columns))
                write_hierarchical_headers(sheet, titles, start_row=plan.header_row)
            column_widths = [max(column.width, 10, max(display_width(part) for part in _header_paths([column.title])[0]) + 2)
                             for column in plan.columns]
            for col_index, column in enumerate(plan.columns, start=1):
                sheet.column_dimensions[get_column_letter(col_index)].width = column_widths[col_index - 1]
            if not reference:
                for row in range(plan.header_row, data_start_row):
                    sheet.row_dimensions[row].height = 26
            written = 0
            for record in records_group:
                dimensions = [[value.strip() for value in record["values_json"].get(field, "").split(plan.sku_separator)] for field in plan.sku_fields]
                combinations = itertools.product(*dimensions) if dimensions else [()]
                for combination in combinations:
                    written += 1
                    if written > MAX_ROWS:
                        raise WorkflowError("SKU 展开或输出记录超过每表 20,000 条限制。")
                    values = {**record["values_json"], **dict(zip(plan.sku_fields, combination))}
                    index = data_start_row + written - 1
                    for col_index, column in enumerate(plan.columns, start=1):
                        cell = sheet.cell(index, col_index)
                        cell._style = copy(styles[col_index - 1])
                        if column.operation == "formula":
                            cell.value = column.argument.replace("{row}", str(index))
                        else:
                            sources = column.sources if column.operation == "concat" else [column.field]
                            value = transform([values.get(field, "") for field in sources], column.operation, column.argument, column.part)
                            if column.required and not value.strip():
                                raise WorkflowError(f"{name} 第 {written} 条记录缺少输出必填字段：{column.title}")
                            rule = rules.get(column.field) if column.operation in {"copy", "multiply"} else None
                            typed_cell(cell, value, rule)
                            column_widths[col_index - 1] = min(60, max(column_widths[col_index - 1], display_width(value) + 2))
                        if not reference:
                            style_data_cell(cell)
                        if column.number_format:
                            cell.number_format = column.number_format
            for col_index, width in enumerate(column_widths, start=1):
                sheet.column_dimensions[get_column_letter(col_index)].width = width
            last = max(data_start_row, data_start_row + written - 1)
            sheet.freeze_panes = f"A{data_start_row}"
            sheet.auto_filter.ref = f"A{plan.header_row}:{get_column_letter(len(plan.columns))}{last}"
            for col_index, column in enumerate(plan.columns, start=1):
                rule = rules.get(column.field)
                if rule and rule.enum_values:
                    joined = ",".join(rule.enum_values)
                    if len(joined) <= 250 and not any('"' in v or "," in v for v in rule.enum_values):
                        validation = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=not rule.required)
                        validation.error = "请选择模板允许的值"
                        validation.showErrorMessage = True
                        sheet.add_data_validation(validation)
                        letter = get_column_letter(col_index)
                        validation.add(f"{letter}{data_start_row}:{letter}{last}")
            stats.append({"sheet": name, "records": written, "columns": len(plan.columns)})
    for existing in list(book.worksheets):
        if existing.sheet_state == "visible" and existing.title.casefold() not in seen_names:
            book.remove(existing)
    book.save(path)
    book.close()
    verified = load_workbook(path, read_only=True, data_only=False)
    try:
        if any(stat["sheet"] not in verified.sheetnames for stat in stats):
            raise WorkflowError("输出工作簿回读校验失败。")
    finally:
        verified.close()
    return {"name": path.name, "sheets": stats, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "size": path.stat().st_size}
