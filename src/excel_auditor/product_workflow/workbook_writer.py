from __future__ import annotations

import json
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .fixed_template import FIXED_FIELDS
from .merchant_models import ReconciliationIssue, ReconciliationResult


_DIFFERENCE_FILL = "FDE8E7"
_DIFFERENCE_TEXT = "B42318"
_GRID = "000000"
_TEXT = "000000"
_FONT = "Microsoft YaHei"

_FIELD_WIDTHS = {
    "product_category": 12,
    "is_bundle": 12,
    "product_name": 24,
    "sales_unit": 9,
    "quantity_mode": 10,
    "brand": 12,
    "main_image": 16,
    "detail_images": 18,
    "sku": 18,
    "spec_image": 16,
    "market_price": 10,
    "sale_price": 10,
    "stock": 8,
    "handling_fee_template": 14,
    "package_quantity": 9,
    "package_unit": 9,
    "return_days": 10,
    "service_description": 18,
    "warranty_days": 10,
    "after_sales_description": 18,
    "delivery_method": 12,
    "freight_template": 14,
    "home_entry_fee_included": 20,
    "other_description": 20,
    "listed_at": 16,
}


def write_reconciliation_workbook(path: Path, result: ReconciliationResult) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "商品汇总"
    issues = workbook.create_sheet("问题清单")
    trace = workbook.create_sheet("_来源追踪")
    trace.sheet_state = "hidden"

    platform_fields = list(dict.fromkeys(result.platform_extension_fields))
    total_columns = len(FIXED_FIELDS) + len(platform_fields)
    if total_columns > 16_384:
        raise ValueError("OUTPUT_LIMIT_EXCEEDED: too many output fields for an Excel worksheet")

    column_by_field = _write_summary_headers(summary, platform_fields)
    _style_summary(summary, total_columns, len(result.output_records) + 3, column_by_field)
    mismatch_lookup = _mismatch_lookup(result.issues)
    difference_fill = PatternFill("solid", fgColor=_DIFFERENCE_FILL)
    difference_font = Font(name=_FONT, size=10, color=_DIFFERENCE_TEXT)
    for output_index, record in enumerate(result.output_records, start=1):
        excel_row = output_index + 3
        for fixed_field in FIXED_FIELDS:
            cell = summary.cell(excel_row, column_by_field[fixed_field.field_id])
            cell.value = _cell_value(record.fixed.get(fixed_field.field_id))
            if fixed_field.field_id in record.difference_fields:
                cell.fill = difference_fill
                cell.font = difference_font
                issue = mismatch_lookup.get((record.source_sheet, record.source_row, fixed_field.field_id))
                cell.comment = Comment(_difference_comment(record.row_status, issue), "ExcelFlow")
        for title in platform_fields:
            summary.cell(excel_row, column_by_field[f"platform:{title}"], _cell_value(record.platform_extras.get(title)))

    _write_issues_sheet(issues, result.issues)
    _write_trace_sheet(trace, result)
    workbook.properties.title = "商家商品整理结果"
    workbook.properties.subject = "固定字段整理、平台对比和来源追踪"
    workbook.properties.creator = "ExcelFlow"
    temporary = path.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    temporary.replace(path)


def _write_summary_headers(sheet: Any, platform_fields: list[str]) -> dict[str, int]:
    columns: dict[str, int] = {}
    fixed_start = 1
    fixed_end = fixed_start + len(FIXED_FIELDS) - 1
    sheet.merge_cells(start_row=1, start_column=fixed_start, end_row=1, end_column=fixed_end)
    sheet.cell(1, fixed_start, "固定字段")

    column = fixed_start
    index = 0
    while index < len(FIXED_FIELDS):
        field = FIXED_FIELDS[index]
        columns[field.field_id] = column
        if field.group is None:
            sheet.merge_cells(start_row=2, start_column=column, end_row=3, end_column=column)
            sheet.cell(2, column, field.title)
            column += 1
            index += 1
            continue
        group = field.group
        group_start = column
        while index < len(FIXED_FIELDS) and FIXED_FIELDS[index].group == group:
            grouped = FIXED_FIELDS[index]
            columns[grouped.field_id] = column
            sheet.cell(3, column, grouped.title)
            column += 1
            index += 1
        sheet.merge_cells(start_row=2, start_column=group_start, end_row=2, end_column=column - 1)
        sheet.cell(2, group_start, group)

    _write_extension_headers(sheet, column, platform_fields, "动态字段（运营平台提供接口）", "platform", columns)
    return columns


def _write_extension_headers(
    sheet: Any,
    start_column: int,
    fields: list[str],
    group_title: str,
    prefix: str,
    columns: dict[str, int],
) -> int:
    if not fields:
        return start_column
    end_column = start_column + len(fields) - 1
    sheet.merge_cells(start_row=1, start_column=start_column, end_row=1, end_column=end_column)
    group_cell = sheet.cell(1, start_column, group_title)
    for offset, title in enumerate(fields):
        column = start_column + offset
        columns[f"{prefix}:{title}"] = column
        sheet.merge_cells(start_row=2, start_column=column, end_row=3, end_column=column)
        sheet.cell(2, column, title)
    return end_column + 1


def _style_summary(
    sheet: Any,
    total_columns: int,
    last_row: int,
    column_by_field: dict[str, int],
) -> None:
    thin = Side(style="thin", color=_GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white = PatternFill("solid", fgColor="FFFFFF")
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 100
    # Keep the fixed-field table as one continuous view, without a frozen split.
    sheet.freeze_panes = None
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:3"
    for row in range(1, 4):
        for column in range(1, total_columns + 1):
            cell = sheet.cell(row, column)
            cell.fill = white
            cell.font = Font(name=_FONT, size=10, bold=True, color=_TEXT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 30
    sheet.row_dimensions[3].height = 30
    for row in range(4, last_row + 1):
        sheet.row_dimensions[row].height = 32
        for column in range(1, total_columns + 1):
            cell = sheet.cell(row, column)
            cell.fill = white
            cell.font = Font(name=_FONT, size=10, color=_TEXT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
    price_fields = {"market_price", "sale_price"}
    numeric_fields = {"stock", "package_quantity", "return_days", "warranty_days"}
    for field_id, column in column_by_field.items():
        sheet.column_dimensions[get_column_letter(column)].width = _FIELD_WIDTHS.get(field_id, 12)
        if field_id.startswith("platform:"):
            continue
        if field_id in price_fields | numeric_fields:
            for row in range(4, last_row + 1):
                sheet.cell(row, column).alignment = Alignment(horizontal="right", vertical="center")
                sheet.cell(row, column).number_format = "#,##0.00" if field_id in price_fields else "#,##0.##"


def _write_issues_sheet(sheet: Any, issues: list[ReconciliationIssue]) -> None:
    headers = ["问题类型", "严重程度", "来源工作表", "来源行", "字段", "商家值", "平台值", "平台记录ID", "匹配分数", "说明"]
    sheet.append(headers)
    for issue in issues:
        sheet.append([
            issue.issue_type,
            issue.severity,
            issue.source_sheet,
            issue.source_row,
            issue.field_title or issue.field_id,
            _cell_value(issue.merchant_value),
            _cell_value(issue.platform_value),
            issue.platform_record_id,
            issue.match_score,
            issue.message,
        ])
    _style_support_sheet(sheet, len(headers), max(1, len(issues) + 1))


def _write_trace_sheet(sheet: Any, result: ReconciliationResult) -> None:
    headers = ["输出行", "来源工作表", "来源行", "平台记录ID", "匹配分数", "状态"]
    sheet.append(headers)
    for index, record in enumerate(result.output_records, start=4):
        sheet.append([
            index,
            record.source_sheet,
            record.source_row,
            record.platform_record_id,
            record.match_score,
            record.row_status,
        ])
    _style_support_sheet(sheet, len(headers), max(1, len(result.output_records) + 1))


def _style_support_sheet(sheet: Any, column_count: int, row_count: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 100
    sheet.freeze_panes = None
    sheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}{row_count}"
    thin = Side(style="thin", color=_GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white = PatternFill("solid", fgColor="FFFFFF")
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.fill = white
        cell.font = Font(name=_FONT, size=10, bold=True, color=_TEXT)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    widths = [20, 18, 10, 20, 12, 20, 20, 22, 12, 52]
    for index in range(1, column_count + 1):
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    for row in sheet.iter_rows(min_row=2, max_row=row_count):
        for cell in row:
            cell.fill = white
            cell.font = Font(name=_FONT, size=10, color=_TEXT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        sheet.row_dimensions[row[0].row].height = 30


def _mismatch_lookup(issues: list[ReconciliationIssue]) -> dict[tuple[str | None, int | None, str | None], ReconciliationIssue]:
    return {
        (issue.source_sheet, issue.source_row, issue.field_id): issue
        for issue in issues
        if issue.issue_type == "field_mismatch"
    }


def _difference_comment(status: str, issue: ReconciliationIssue | None) -> str:
    if issue is not None:
        lines = (
            f"商家值：{_display_value(issue.merchant_value)}\n"
            f"平台值：{_display_value(issue.platform_value)}\n"
            f"来源：{issue.source_sheet}!{issue.source_row}\n"
        )
        if issue.match_score is not None:
            lines += f"匹配分数：{issue.match_score:g}"
        return lines[:30_000]
    if status == "ambiguous":
        return "存在多个同分平台候选，未自动合并。"
    return "未找到相似度达到阈值的平台商品。"


def _cell_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.isoformat()
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    if isinstance(value, str):
        if len(value) > 32_767:
            raise ValueError("OUTPUT_VALUE_TOO_LONG: a cell value exceeds Excel's text limit")
        if ILLEGAL_CHARACTERS_RE.search(value):
            raise ValueError("OUTPUT_VALUE_INVALID: a cell contains characters Excel cannot store")
        if value.startswith("="):
            return "'" + value
    return value


def _display_value(value: Any) -> str:
    if value is None:
        return "（空）"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def normalize_header_text(value: Any) -> str:
    return "" if value is None else str(value).replace("\n", "").replace(" ", "")
