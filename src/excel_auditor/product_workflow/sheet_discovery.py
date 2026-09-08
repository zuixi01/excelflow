from __future__ import annotations

from collections.abc import Iterable
from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from .fixed_template import match_fixed_field, normalize_header, normalize_text
from .merchant_models import DiscoveredSheet


_AUXILIARY_FIELDS = {
    "product_category",
    "sales_unit",
    "brand",
    "sku",
    "market_price",
    "sale_price",
}


def logical_bounds(sheet: Worksheet) -> tuple[int, int]:
    """Return the last row/column containing a value, ignoring format-only cells."""
    last_row = 0
    last_column = 0
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.value not in {None, ""}:
                last_row = max(last_row, cell.row)
                last_column = max(last_column, cell.column)
    return last_row, last_column


def merged_value_map(sheet: Worksheet, *, last_row: int, last_column: int) -> dict[tuple[int, int], Any]:
    values: dict[tuple[int, int], Any] = {}
    for merged in sheet.merged_cells.ranges:
        if merged.min_row > last_row or merged.min_col > last_column:
            continue
        value = sheet.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row, min(merged.max_row, last_row) + 1):
            for column in range(merged.min_col, min(merged.max_col, last_column) + 1):
                values[(row, column)] = value
    return values


def discover_product_sheets(workbook: Any, *, scan_rows: int = 20) -> list[DiscoveredSheet]:
    discovered: list[DiscoveredSheet] = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        last_row, last_column = logical_bounds(sheet)
        if last_row < 2 or last_column < 2 or _is_output_template(sheet, last_row, last_column):
            continue
        merged = merged_value_map(sheet, last_row=last_row, last_column=last_column)
        candidate = _best_header_row(
            sheet,
            merged,
            last_row=min(last_row, scan_rows),
            last_column=last_column,
        )
        if candidate is None:
            continue
        header_row, score = candidate
        header_depth = _header_depth(sheet, header_row, last_row, last_column)
        headers = flatten_headers(
            sheet,
            merged,
            header_row=header_row,
            header_depth=header_depth,
            last_column=last_column,
        )
        discovered.append(DiscoveredSheet(
            name=sheet.title,
            header_row=header_row,
            header_depth=header_depth,
            headers=headers,
            last_row=last_row,
            score=score,
        ))
    return discovered


def flatten_headers(
    sheet: Worksheet,
    merged: dict[tuple[int, int], Any],
    *,
    header_row: int,
    header_depth: int,
    last_column: int,
) -> list[str]:
    headers: list[str] = []
    for column in range(1, last_column + 1):
        parts: list[str] = []
        for row in range(header_row, header_row + header_depth):
            raw = merged.get((row, column), sheet.cell(row, column).value)
            text = normalize_text(raw)
            if text and (not parts or normalize_header(parts[-1]) != normalize_header(text)):
                parts.append(text)
        headers.append(" / ".join(parts))
    return headers


def _best_header_row(
    sheet: Worksheet,
    merged: dict[tuple[int, int], Any],
    *,
    last_row: int,
    last_column: int,
) -> tuple[int, int] | None:
    best: tuple[int, int] | None = None
    for row in range(1, last_row + 1):
        field_ids = {
            field_id
            for column in range(1, last_column + 1)
            if (field_id := match_fixed_field(merged.get((row, column), sheet.cell(row, column).value)))
        }
        if "product_name" not in field_ids or len(field_ids & _AUXILIARY_FIELDS) < 2:
            continue
        nonempty = sum(
            1
            for column in range(1, last_column + 1)
            if normalize_text(merged.get((row, column), sheet.cell(row, column).value))
        )
        score = len(field_ids) * 100 + min(nonempty, 99)
        if best is None or score > best[1]:
            best = (row, score)
    return best


def _header_depth(sheet: Worksheet, header_row: int, last_row: int, last_column: int) -> int:
    depth = 1
    for merged in sheet.merged_cells.ranges:
        if merged.min_row != header_row or merged.min_col > last_column:
            continue
        if merged.max_row > header_row:
            depth = max(depth, min(3, merged.max_row - header_row + 1))
        elif merged.max_col > merged.min_col and header_row < last_row:
            child_values = [
                normalize_text(sheet.cell(header_row + 1, column).value)
                for column in range(merged.min_col, min(merged.max_col, last_column) + 1)
            ]
            if sum(bool(value) for value in child_values) >= 2:
                depth = max(depth, 2)
    return min(depth, max(1, last_row - header_row))


def _is_output_template(sheet: Worksheet, last_row: int, last_column: int) -> bool:
    markers = {
        normalize_header(sheet.cell(row, column).value)
        for row in range(1, min(last_row, 4) + 1)
        for column in range(1, min(last_column, 80) + 1)
        if sheet.cell(row, column).value not in {None, ""}
    }
    return "固定字段" in markers and any(marker.startswith("动态字段") for marker in markers)


def ordered_unique(values: Iterable[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))
