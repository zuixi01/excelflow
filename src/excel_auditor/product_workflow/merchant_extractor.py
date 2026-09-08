from __future__ import annotations

import re
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from .fixed_template import FIXED_FIELDS, match_fixed_field, normalize_text
from .merchant_models import DiscoveredSheet, MerchantRecord
from .sheet_discovery import merged_value_map


def extract_merchant_records(
    workbook: Any,
    discovered_sheets: list[DiscoveredSheet],
) -> list[MerchantRecord]:
    records: list[MerchantRecord] = []
    for sheet_index, discovery in enumerate(discovered_sheets, start=1):
        sheet = workbook[discovery.name]
        last_column = len(discovery.headers)
        merged = merged_value_map(sheet, last_row=discovery.last_row, last_column=last_column)
        fixed_columns = _fixed_column_plan(discovery.headers)
        data_start = discovery.header_row + discovery.header_depth
        for row in range(data_start, discovery.last_row + 1):
            if _is_note_or_empty_row(sheet, merged, row, last_column):
                continue
            fixed = {field.field_id: None for field in FIXED_FIELDS}
            for field_id, column in fixed_columns.items():
                fixed[field_id] = _effective_value(sheet, merged, row, column)
            if not normalize_text(fixed.get("product_name")):
                continue
            records.append(MerchantRecord(
                record_id=f"merchant-{sheet_index}-{row}",
                source_sheet=sheet.title,
                source_row=row,
                fixed=fixed,
            ))
    return records


def _fixed_column_plan(headers: list[str]) -> dict[str, int]:
    fixed_columns: dict[str, int] = {}
    for column, header in enumerate(headers, start=1):
        if not normalize_text(header):
            continue
        field_id = match_fixed_field(header)
        if field_id and field_id not in fixed_columns:
            fixed_columns[field_id] = column
    return fixed_columns


def _effective_value(
    sheet: Worksheet,
    merged: dict[tuple[int, int], Any],
    row: int,
    column: int,
) -> Any:
    return merged.get((row, column), sheet.cell(row, column).value)


def _is_note_or_empty_row(
    sheet: Worksheet,
    merged: dict[tuple[int, int], Any],
    row: int,
    last_column: int,
) -> bool:
    raw_values = [sheet.cell(row, column).value for column in range(1, last_column + 1)]
    nonempty = [value for value in raw_values if value not in {None, ""}]
    if not nonempty:
        return True
    first_text = normalize_text(nonempty[0])
    if len(nonempty) == 1 and re.match(r"^(备注|说明|注[:：]?)", first_text):
        return True
    if len(nonempty) == 1:
        for region in sheet.merged_cells.ranges:
            if region.min_row <= row <= region.max_row and region.max_col - region.min_col + 1 >= max(4, last_column // 2):
                return True
    return False
