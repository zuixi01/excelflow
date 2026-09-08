from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from excel_auditor.product_workflow.fixed_template import FIXED_FIELDS
from excel_auditor.product_workflow.merchant_extractor import extract_merchant_records
from excel_auditor.product_workflow.merchant_models import MerchantRecord, PlatformRecord
from excel_auditor.product_workflow.merchant_service import MerchantReconciliationService
from excel_auditor.product_workflow.reconciler import reconcile_products
from excel_auditor.product_workflow.sheet_discovery import discover_product_sheets
from excel_auditor.product_workflow.workbook_writer import write_reconciliation_workbook
from excel_auditor.service import AuditService


def test_discovers_three_merchant_layouts_and_extracts_77_records(tmp_path: Path):
    source = tmp_path / "merchant.xlsx"
    _build_multi_sheet_workbook(source)
    workbook = load_workbook(source, data_only=False, keep_links=False)

    discovered = discover_product_sheets(workbook)
    records = extract_merchant_records(workbook, discovered)

    assert [(sheet.name, sheet.header_row, sheet.header_depth) for sheet in discovered] == [
        ("封窗", 1, 1),
        ("瓷砖", 1, 2),
        ("灯具", 2, 1),
    ]
    assert {name: sum(record.source_sheet == name for record in records) for name in ("封窗", "瓷砖", "灯具")} == {
        "封窗": 8,
        "瓷砖": 37,
        "灯具": 32,
    }
    assert len(records) == 77
    assert all(set(record.fixed) == {field.field_id for field in FIXED_FIELDS} for record in records)
    assert all(record.fixed["product_category"] is None for record in records)
    workbook.close()


def test_reconciliation_matches_fills_marks_differences_and_appends(tmp_path: Path):
    merchant = [
        MerchantRecord(
            record_id="m-1",
            source_sheet="灯具",
            source_row=3,
            fixed=_fixed(product_name="三雄极光 防眩筒灯 3 寸", brand="三雄极光", sale_price=10),
        ),
        MerchantRecord(
            record_id="m-2",
            source_sheet="灯具",
            source_row=4,
            fixed=_fixed(product_name="商家独有商品", brand="测试品牌"),
        ),
    ]
    platform = [
        PlatformRecord(
            record_id="p-1",
            fixed=_fixed(product_name="三雄极光防眩筒灯3寸", brand="三雄极光", market_price=15, sale_price=12),
            extras={"电压": "220V"},
        ),
        PlatformRecord(
            record_id="p-2",
            fixed=_fixed(product_name="平台新增商品", brand="平台品牌"),
            extras={"电压": "110V"},
        ),
    ]

    result = reconcile_products(
        [], merchant, platform, ["电压"], threshold=90, platform_available=True,
    )

    assert len(result.matches) == 1
    assert result.matches[0].score == 100
    assert len(result.output_records) == 3
    assert result.output_records[0].fixed["market_price"] == 15
    assert result.output_records[0].fixed["sale_price"] == 10
    assert result.output_records[0].difference_fields == {"sale_price"}
    assert result.output_records[1].difference_fields == set()
    assert result.merchant_only_count == 1
    assert result.platform_only_count == 1
    assert {issue.issue_type for issue in result.issues} == {
        "field_mismatch",
        "merchant_only_product",
        "platform_only_product",
    }

    output = tmp_path / "result.xlsx"
    write_reconciliation_workbook(output, result)
    rendered = load_workbook(output, data_only=False)
    assert rendered.sheetnames == ["商品汇总", "问题清单", "_来源追踪"]
    assert all(rendered[name].freeze_panes is None for name in rendered.sheetnames)
    assert rendered["_来源追踪"].sheet_state == "hidden"
    sale_price_column = 1 + next(index for index, field in enumerate(FIXED_FIELDS) if field.field_id == "sale_price")
    market_price_column = 1 + next(index for index, field in enumerate(FIXED_FIELDS) if field.field_id == "market_price")
    assert rendered["商品汇总"].cell(4, market_price_column).value == 15
    mismatch = rendered["商品汇总"].cell(4, sale_price_column)
    assert mismatch.value == 10
    assert mismatch.fill.fgColor.rgb.endswith("FDE8E7")
    assert mismatch.comment is not None and "平台值：12" in mismatch.comment.text
    assert rendered["商品汇总"].cell(6, 3).value == "平台新增商品"
    header_values = {
        rendered["商品汇总"].cell(row, column).value
        for row in range(1, 4)
        for column in range(1, rendered["商品汇总"].max_column + 1)
    }
    assert "商家扩展字段" not in header_values
    assert "内部备注" not in header_values
    assert "表头" not in header_values
    rendered.close()


def test_api_creates_new_workbook_with_trace_and_preserves_staged_input(tmp_path: Path, monkeypatch):
    import excel_auditor.api as api_module

    audit_service = AuditService(tmp_path / "runtime")
    monkeypatch.setattr(api_module, "service", audit_service)
    monkeypatch.setattr(api_module, "merchant_reconciliation_service", MerchantReconciliationService(audit_service))
    monkeypatch.setattr(api_module, "task_queue", None)
    monkeypatch.setattr(api_module, "database", None)
    monkeypatch.setattr(api_module, "artifact_store", None)
    content = _minimal_workbook_bytes()
    original_hash = hashlib.sha256(content).hexdigest()
    fixture = {
        "records": [{
            "platform_id": "platform-1",
            "商品名称": "测试灯具A",
            "商品品牌": "测试品牌",
            "销售价": 89,
            "dynamic_fields": {"电压": "220V"},
        }]
    }

    client = TestClient(api_module.app)
    response = client.post(
        "/api/v1/merchant-product-reconciliations",
        files={"excel_file": ("merchant.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"platform_json": json.dumps(fixture, ensure_ascii=False), "name_match_threshold": "90"},
        headers={"Idempotency-Key": "merchant-flow-test"},
    )

    assert response.status_code == 202
    job_id = response.json()["job_id"]
    status = client.get(f"/api/v1/comparisons/{job_id}").json()
    assert status["status"] == "completed"
    assert status["summary"] == {
        "recognized_sheets": 1,
        "source_records": 1,
        "platform_records": 1,
        "matched_records": 1,
        "differences": 1,
        "merchant_only": 0,
        "platform_only": 0,
        "output_records": 1,
    }
    staged = audit_service.job_directory(job_id) / "merchant-input.xlsx"
    assert hashlib.sha256(staged.read_bytes()).hexdigest() == original_hash

    download = client.get(f"/api/v1/comparisons/{job_id}/artifacts/product_excel")
    assert download.status_code == 200
    result_book = load_workbook(io.BytesIO(download.content), data_only=False)
    assert result_book["商品汇总"]["A1"].value == "固定字段"
    assert result_book["商品汇总"]["A2"].value == "商品分类"
    assert result_book["商品汇总"].freeze_panes is None
    assert result_book["问题清单"].max_row == 2
    assert result_book["_来源追踪"].sheet_state == "hidden"
    result_book.close()

    issues = client.get(f"/api/v1/merchant-product-reconciliations/{job_id}/issues").json()
    assert issues["total"] == 1
    assert issues["items"][0]["issue_type"] == "field_mismatch"


def _fixed(**values):
    fixed = {field.field_id: None for field in FIXED_FIELDS}
    fixed.update(values)
    return fixed


def _minimal_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品"
    sheet.append(["商品名称", "品牌", "计价单位", "类别", "市场价", "销售价"])
    sheet.append(["测试灯具 A", "测试品牌", "个", "灯具", 100, 88])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _build_multi_sheet_workbook(path: Path) -> None:
    workbook = Workbook()
    windows = workbook.active
    windows.title = "封窗"
    windows.append(["序号", "商品名称", "品牌", "计价单位", "类别", "市场价"])
    for index in range(4):
        start = 2 + index * 2
        windows.cell(start, 1, index + 1)
        windows.cell(start, 2, f"窗商品{index + 1}")
        windows.merge_cells(start_row=start, start_column=2, end_row=start + 1, end_column=2)
        windows.cell(start, 3, "窗品牌")
        windows.merge_cells(start_row=start, start_column=3, end_row=start + 1, end_column=3)
        for row in (start, start + 1):
            windows.cell(row, 4, "平方米")
            windows.cell(row, 5, "固定位置" if row == start else "开启扇")
            windows.cell(row, 6, 100 + row)

    tiles = workbook.create_sheet("瓷砖")
    tiles.merge_cells("A1:A2"); tiles["A1"] = "序号"
    tiles.merge_cells("B1:B2"); tiles["B1"] = "品牌"
    tiles.merge_cells("C1:C2"); tiles["C1"] = "商品名称"
    tiles.merge_cells("D1:D2"); tiles["D1"] = "型号"
    tiles.merge_cells("E1:F1"); tiles["E1"] = "属性"
    tiles["E2"] = "产地"; tiles["F2"] = "厚度"
    tiles.merge_cells("G1:G2"); tiles["G1"] = "计价单位"
    tiles.merge_cells("H1:H2"); tiles["H1"] = "市场价"
    for index, row in enumerate(range(3, 40), start=1):
        tiles.append([index, "瓷砖品牌", f"瓷砖商品{index}", f"T{index:03d}", "长沙", "8mm", "片", 20])

    lamps = workbook.create_sheet("灯具")
    lamps.merge_cells("A1:H1"); lamps["A1"] = "灯具价格清单"
    lamps.append(["产品品牌", "产品名称", "产品型号", "元素(图片）", "基本单位", "市场价", "建议平台卖价", "材质"])
    for index in range(1, 33):
        lamps.append(["灯具品牌", f"灯具商品{index}", f"L{index:03d}", None, "个", 100, 80, "铝"])
    lamps.merge_cells("A35:H35"); lamps["A35"] = "备注：价格仅供参考"
    lamps["A200"].fill = PatternFill("solid", fgColor="FFFFFF")

    template = workbook.create_sheet("Sheet2")
    template["B1"] = "固定字段"
    template["AA1"] = "动态字段（运营平台提供接口）"
    template["A2"] = "表头"
    template["B2"] = "商品分类"
    template["D2"] = "商品名称"
    workbook.save(path)
    workbook.close()
