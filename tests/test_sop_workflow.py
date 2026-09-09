from __future__ import annotations

import importlib
import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from pydantic import SecretStr, ValidationError
from sqlalchemy import select

from excel_auditor.workflow.ai import AIClient, AISettings
from excel_auditor.workflow.excel import discover, safe_book, write_output
from excel_auditor.workflow.models import EditRequest, MappingRequest, OutputColumn, TemplateEdit, WorkflowConfig, WorkflowError, default_config
from excel_auditor.workflow.repository import WorkflowStore, objects, records, revisions
from excel_auditor.workflow.tasks import TaskService
from excel_auditor.workflow.templates import TemplateService


def workbook(rows=None, headers=None, title="商家商品"):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers or ["商品名称", "货号", "销售价", "库存", "商家备注"])
    for row in (rows if rows is not None else [["台灯", "0001", "29.50", "8", "暖光"], ["吊灯", "0002", "199.90", "3", "客厅"]]):
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


@pytest.fixture
def runtime(tmp_path):
    store = WorkflowStore(tmp_path / "workflow")
    yield store, TaskService(store), TemplateService(store)
    store.engine.dispose()


def create_task(runtime, content=None, tenant="local", template_id=None):
    store, service, templates = runtime
    if template_id is None:
        draft = templates.create(tenant, name="测试流程")
        configured = templates.update(draft["id"], tenant, TemplateEdit(
            base_revision=draft["revision"], config=default_config(), confirmed=True,
        ))
        template_id = templates.publish(configured["id"], tenant, configured["revision"])["id"]
    task, _ = service.create(tenant, "local", template_id, "测试整理", "input.xlsx", content or workbook())
    service.parse(task["id"], tenant)
    return service.get(task["id"], tenant)


def mapping(task):
    return MappingRequest(base_revision=task["revision"], sheets=[{
        "sheet": sheet["sheet"], "header_row": sheet["header_row"], "header_depth": sheet["header_depth"],
        "bindings": sheet["suggested_bindings"],
    } for sheet in task["payload"]["sheets"]])


def ready_task(runtime, **kwargs):
    task = create_task(runtime, **kwargs)
    assert task["status"] == "mapping", task
    return runtime[1].confirm_mapping(task["id"], kwargs.get("tenant", "local"), "local", mapping(task))


def test_maintenance_persists_exact_text_and_cannot_overwrite_new_revision(runtime):
    store, service, _ = runtime
    task = ready_task(runtime)
    assert task["status"] == "ready"
    rows = service.rows(task["id"], "local", 0, 30)["items"]
    assert rows[0]["values_json"]["sku"] == "0001"
    assert rows[0]["values_json"]["sale_price"] == "29.50"
    assert not any(key.startswith("extra_") for key in rows[0]["values_json"])
    updated = service.edit(task["id"], "local", "local", EditRequest(base_revision=1, changes={rows[0]["id"]: {"sale_price": "31.60"}}))
    assert updated["revision"] == 2
    with pytest.raises(WorkflowError, match="新版本"):
        service.edit(task["id"], "local", "local", EditRequest(base_revision=1, changes={rows[0]["id"]: {"sale_price": "1"}}))
    second = WorkflowStore(store.root)
    try:
        reread = TaskService(second).rows(task["id"], "local", 0, 30)["items"][0]
        assert reread["values_json"]["sale_price"] == "31.60"
        assert reread["original"]["sale_price"] == "29.50"
        assert reread["source"] == {"sheet": "商家商品", "row": 2}
    finally:
        second.engine.dispose()


def test_required_and_type_errors_block_export_and_are_fixable(runtime):
    _, service, _ = runtime
    task = ready_task(runtime, content=workbook([["", "0001", "bad", "5", "备注"], ["台灯", "0002", "30", "5", ""]]))
    assert task["status"] == "maintenance"
    with pytest.raises(WorkflowError, match="阻断"):
        service.create_export(task["id"], "local", 1)
    rows = service.rows(task["id"], "local", 0, 20, issues_only=True)
    assert rows["total"] == 1
    fixed = service.edit(task["id"], "local", "local", EditRequest(base_revision=1, changes={rows["items"][0]["id"]: {"product_name": "壁灯", "sale_price": "30"}}))
    assert fixed["status"] == "ready" and fixed["payload"]["error_count"] == 0


def test_duplicate_mapping_does_not_advance_task(runtime):
    _, service, _ = runtime
    task = create_task(runtime, workbook([["灯", "A", "001", "3"]], ["商品名称", "货号", "型号", "销售价"]))
    with pytest.raises(WorkflowError, match="同一字段"):
        service.confirm_mapping(task["id"], "local", "local", mapping(task))
    assert service.get(task["id"], "local")["revision"] == 0


def test_maintenance_template_uses_actual_multilevel_headers_not_presets(runtime, tmp_path):
    _, _, templates = runtime
    path = tmp_path / "sparse-template.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "基础资料"
    sheet["C1"] = "扩展资料"
    sheet.append(["商家编码", "自定义名称", "接口属性"])
    # This is a common Excel artifact: a blank but formatted far-right cell.
    sheet.cell(2, 16369).fill = PatternFill("solid", fgColor="FFFFFF")
    book.save(path)
    book.close()

    sheets = discover(path, default_config().fields, template=True)
    assert [(sheet["sheet"], len(sheet["columns"])) for sheet in sheets] == [("Sheet", 3)]

    template = templates.upload("local", "标准维护", path.read_bytes(), "maintenance")
    templates.parse(template["id"], "local")
    parsed = templates.get(template["id"], "local")
    assert parsed["status"] == "draft"
    fields = parsed["payload"]["config"]["fields"]
    assert [field["rule"]["title"] for field in fields] == ["基础资料 / 商家编码", "基础资料 / 自定义名称", "扩展资料 / 接口属性"]
    assert all(field["rule"]["name"].startswith("field_") for field in fields)
    assert all(field["rule"]["type"] == "string" and field["editable"] for field in fields)
    assert not parsed["payload"]["dynamic_columns"]

    output = templates.attach_source(template["id"], "local", workbook(rows=[], headers=["商家编码", "接口属性"]), "output")
    templates.parse(output["id"], "local")
    combined = templates.get(template["id"], "local")
    assert len(combined["payload"]["config"]["fields"]) == 3
    assert [column["title"] for column in combined["payload"]["config"]["outputs"][0]["sheets"][0]["columns"]] == ["商家编码", "接口属性"]
    assert set(combined["payload"]["sources"]) == {"maintenance", "output"}


def test_reopen_mapping_clears_maintenance_rows_and_keeps_mapping_snapshot(runtime):
    _, service, _ = runtime
    task = ready_task(runtime)
    reopened = service.reopen_mapping(task["id"], "local", "local", task["revision"])
    assert reopened["status"] == "mapping"
    assert reopened["payload"]["record_count"] == 0
    assert reopened["payload"]["mapping"]["sheets"]
    assert service.rows(task["id"], "local", 0, 20)["total"] == 0


def test_roundtrip_sort_add_and_missing_rows_default_to_preservation(runtime):
    _, service, _ = runtime
    task = ready_task(runtime)
    path = service.maintenance_file(task["id"], "local")
    book = load_workbook(path)
    sheet = book["维护表"]
    headers = ["_workflow_record_id", *(field["rule"]["title"] for field in task["payload"]["config"]["fields"])]
    data_start_row = 1 + max(region.max_row for region in sheet.merged_cells.ranges if region.min_row == 1)
    second = [cell.value for cell in sheet[data_start_row + 1]]
    sheet.delete_rows(data_start_row, 2)
    second[headers.index("销售价")] = "219.00"
    sheet.append(second)
    new = [None] * len(headers)
    new[headers.index("商品名称")] = "新增灯"
    new[headers.index("货号")] = "0003"
    sheet.append(new)
    output = BytesIO(); book.save(output); book.close()
    preview = service.preview_return(task["id"], "local", output.getvalue())
    assert preview["payload"]["counts"] == {"changed": 1, "added": 1, "missing": 1}
    updated = service.apply_return(task["id"], "local", "local", preview["id"], False)
    assert updated["payload"]["record_count"] == 3
    assert service.rows(task["id"], "local", 0, 20)["items"][1]["values_json"]["sale_price"] == "219.00"
    with pytest.raises(WorkflowError):
        service.apply_return(task["id"], "local", "local", preview["id"], False)


def test_maintenance_download_restores_grouped_template_headers(runtime):
    _, service, _ = runtime
    task = ready_task(runtime)
    book = load_workbook(service.maintenance_file(task["id"], "local"))
    try:
        sheet = book["维护表"]
        titles = [field["rule"]["title"] for field in task["payload"]["config"]["fields"]]
        package_start = titles.index("包装量/数值") + 2
        package_end = package_start + 1
        assert sheet.cell(1, package_start).value == "包装量"
        assert sheet.cell(2, package_start).value == "数值"
        assert sheet.cell(2, package_end).value == "单位"
        assert f"{get_column_letter(package_start)}1:{get_column_letter(package_end)}1" in {str(region) for region in sheet.merged_cells.ranges}
        assert sheet.freeze_panes == "B3"
        assert sheet.cell(1, package_start).fill.fill_type is None
    finally:
        book.close()


def test_task_delete_removes_records_exports_and_files(runtime):
    store, service, _ = runtime
    task = ready_task(runtime)
    run, _ = service.create_export(task["id"], "local", task["revision"])
    service.generate(run["id"], "local")
    assert service.delete("local", [task["id"]]) == {"deleted": [task["id"]]}
    with pytest.raises(WorkflowError, match="不存在"):
        service.get(task["id"], "local")
    with store.engine.connect() as con:
        assert con.execute(select(records.c.id).where(records.c.task_id == task["id"])).first() is None
        assert con.execute(select(revisions.c.task_id).where(revisions.c.task_id == task["id"])).first() is None
        assert con.execute(select(objects.c.id).where(objects.c.id == run["id"])).first() is None
    assert not (store.root / "files" / task["id"]).exists()
    assert not (store.root / "files" / run["id"]).exists()


@pytest.mark.parametrize("mode", ["stale", "duplicate", "unknown", "foreign", "formula"])
def test_invalid_returns_are_rejected(runtime, mode):
    _, service, _ = runtime
    task = ready_task(runtime)
    path = service.maintenance_file(task["id"], "local")
    book = load_workbook(path)
    data_start_row = 1 + max(region.max_row for region in book["维护表"].merged_cells.ranges if region.min_row == 1)
    if mode == "stale":
        service.edit(task["id"], "local", "local", EditRequest(base_revision=1))
    elif mode == "duplicate": book["维护表"][f"A{data_start_row + 1}"] = book["维护表"][f"A{data_start_row}"].value
    elif mode == "unknown": book["维护表"][f"A{data_start_row}"] = "not-a-record"
    elif mode == "foreign": book["_workflow_meta"]["B1"] = "another-task"
    elif mode == "formula": book["维护表"][f"B{data_start_row}"] = "=1+1"
    output = BytesIO(); book.save(output); book.close()
    with pytest.raises(WorkflowError):
        service.preview_return(task["id"], "local", output.getvalue())


def test_export_is_bound_to_snapshot_and_hidden_ids_do_not_leak(runtime):
    store, service, _ = runtime
    task = ready_task(runtime)
    row = service.rows(task["id"], "local", 0, 20)["items"][0]
    run, created = service.create_export(task["id"], "local", 1, "same-request")
    again, created_again = service.create_export(task["id"], "local", 1, "same-request")
    assert created and not created_again and run["id"] == again["id"]
    service.edit(task["id"], "local", "local", EditRequest(base_revision=1, changes={row["id"]: {"sale_price": "500"}}))
    service.generate(run["id"], "local")
    current = service.get(task["id"], "local")
    assert current["status"] == "ready" and current["exports"][0]["payload"]["revision"] == 1
    file = service.export_file(task["id"], run["id"], "local", "商品导入.xlsx")
    book = load_workbook(file)
    titles = [c.value for c in book.active[1]]
    data_start_row = 1 + max(len(field["rule"]["title"].split("/")) for field in task["payload"]["config"]["fields"])
    assert book.active.cell(data_start_row, titles.index("销售价") + 1).value == 29.5
    assert book.active.cell(data_start_row, titles.index("货号") + 1).value == "0001"
    assert "_workflow_record_id" not in titles
    book.close()
    with ZipFile(service.export_file(task["id"], run["id"], "local", "最终表.zip")) as bundle:
        assert set(bundle.namelist()) == {"商品导入.xlsx", "生成说明.json"}
    assert not current["exports"][0]["payload"]["omitted_fields"]


def test_tenant_isolation_covers_task_template_and_export(runtime):
    _, service, templates = runtime
    task = ready_task(runtime, tenant="alpha")
    assert service.list("beta")["total"] == 0
    with pytest.raises(WorkflowError) as exc:
        service.get(task["id"], "beta")
    assert exc.value.status == 404
    with pytest.raises(WorkflowError): templates.get(task["payload"]["template_id"], "beta")
    with pytest.raises(WorkflowError): service.maintenance_file(task["id"], "beta")


def test_template_publish_requires_confirmation_but_not_a_trial(runtime):
    _, _, templates = runtime
    draft = templates.create("local", name="模板试运行")
    with pytest.raises(WorkflowError): templates.publish(draft["id"], "local", 0)
    confirmed = templates.update(draft["id"], "local", TemplateEdit(base_revision=draft["revision"], config=draft["payload"]["config"], confirmed=True))
    published = templates.publish(draft["id"], "local", confirmed["revision"])
    assert published["status"] == "published"
    with pytest.raises(WorkflowError): templates.update(draft["id"], "local", TemplateEdit(base_revision=published["revision"], config=published["payload"]["config"]))
    fork = templates.create("local", source_id=published["id"])
    assert fork["payload"]["version"] == 2 and fork["status"] == "draft"
    assert not fork["payload"].get("tested_hash")


def test_template_versions_are_presented_as_one_template_with_the_latest_published_definition(runtime):
    _, _, templates = runtime
    original = templates.create("local", name="商品流程")
    confirmed = templates.update(original["id"], "local", TemplateEdit(base_revision=original["revision"], config=original["payload"]["config"], confirmed=True))
    templates.publish(original["id"], "local", confirmed["revision"])
    version_two = templates.create("local", source_id=original["id"])
    assert [item["id"] for item in templates.list("local")["items"] if item["payload"]["config"]["name"] == "商品流程"] == [version_two["id"]]
    confirmed = templates.update(version_two["id"], "local", TemplateEdit(base_revision=version_two["revision"], config=version_two["payload"]["config"], confirmed=True))
    templates.publish(version_two["id"], "local", confirmed["revision"])
    assert [item["id"] for item in templates.definitions("local")["items"] if item["payload"]["config"]["name"] == "商品流程"] == [version_two["id"]]


def test_template_delete_removes_unused_templates_and_preserves_referenced_templates(runtime):
    store, service, templates = runtime
    assert templates.list("local")["items"] == []
    removable = templates.create("local", name="待删除流程")
    template_file = store.directory(removable["id"]) / "reference.xlsx"
    template_file.write_bytes(workbook())
    assert templates.delete("local", [removable["id"]]) == {"deleted": [removable["id"]]}
    with pytest.raises(WorkflowError, match="不存在"):
        templates.get(removable["id"], "local")
    assert not template_file.parent.exists()

    used = templates.create("local", name="已使用流程")
    trial = templates.trial(used["id"], "local", used["revision"], workbook())
    published = templates.update(used["id"], "local", TemplateEdit(base_revision=trial["revision"], config=trial["payload"]["config"], confirmed=True))
    published = templates.publish(used["id"], "local", published["revision"])
    service.create("local", "local", published["id"], "引用流程", "input.xlsx", workbook())
    with pytest.raises(WorkflowError, match="已被任务引用"):
        templates.delete("local", [published["id"]])


def test_changed_template_can_publish_without_a_new_trial(runtime):
    _, _, templates = runtime
    draft = templates.create("local")
    trial = templates.trial(draft["id"], "local", 0, workbook())
    config = WorkflowConfig.model_validate(trial["payload"]["config"])
    config.outputs[0].name = "新输出.xlsx"
    updated = templates.update(draft["id"], "local", TemplateEdit(base_revision=trial["revision"], config=config, confirmed=True))
    assert "tested_hash" not in updated["payload"]
    published = templates.publish(draft["id"], "local", updated["revision"])
    assert published["status"] == "published"


def test_excel_template_structure_can_be_parsed_trialled_and_used(runtime):
    _, service, templates = runtime
    draft = templates.upload("local", "平台模板", workbook(headers=["商品名称", "货号", "销售价"], rows=[["灯", "0001", "10"]], title="平台导入"))
    templates.parse(draft["id"], "local")
    parsed = templates.get(draft["id"], "local")
    assert parsed["status"] == "draft", parsed
    assert parsed["payload"]["config"]["outputs"][0]["sheets"][0]["name"] == "平台导入"
    trial = templates.trial(draft["id"], "local", parsed["revision"])
    confirmed = templates.update(draft["id"], "local", TemplateEdit(base_revision=trial["revision"], config=trial["payload"]["config"], confirmed=True))
    published = templates.publish(draft["id"], "local", confirmed["revision"])
    task = ready_task(runtime, template_id=published["id"])
    run, _ = service.create_export(task["id"], "local", 1)
    service.generate(run["id"], "local")
    current = service.get(task["id"], "local")
    assert current["exports"][0]["status"] == "completed", current


def test_output_template_ignores_structural_header_but_keeps_its_own_fields(runtime):
    _, _, templates = runtime
    book = Workbook()
    sheet = book.active
    sheet.title = "自定义商品表"
    sheet.append(["表头", "商家自定义编码", "特殊规格"])
    sheet.append(["数值", "M-001", "加厚款"])
    content = BytesIO()
    book.save(content)
    book.close()

    draft = templates.upload("local", "不同字段模板", content.getvalue(), "output")
    templates.parse(draft["id"], "local")
    parsed = templates.get(draft["id"], "local")

    assert parsed["status"] == "draft", parsed
    assert [field["rule"]["title"] for field in parsed["payload"]["config"]["fields"]] == ["商家自定义编码", "特殊规格"]
    columns = parsed["payload"]["config"]["outputs"][0]["sheets"][0]["columns"]
    assert [column["title"] for column in columns] == ["商家自定义编码", "特殊规格"]


def test_export_can_use_a_separately_selected_output_template(runtime):
    _, service, templates = runtime
    output = templates.upload("local", "电商平台最终表", workbook(rows=[], headers=["商品名称", "货号", "销售价"], title="平台商品"), "output")
    templates.parse(output["id"], "local")
    parsed = templates.get(output["id"], "local")
    published = templates.update(output["id"], "local", TemplateEdit(
        base_revision=parsed["revision"], config=parsed["payload"]["config"], confirmed=True,
    ))
    published = templates.publish(output["id"], "local", published["revision"])

    task = ready_task(runtime)
    run, created = service.create_export(task["id"], "local", task["revision"], output_template_id=published["id"])
    assert created
    assert run["payload"]["maintenance_template_id"] == task["payload"]["template_id"]
    assert run["payload"]["output_template_id"] == published["id"]
    assert [column["title"] for column in run["payload"]["config"]["outputs"][0]["sheets"][0]["columns"]] == ["商品名称", "货号", "销售价"]

    service.generate(run["id"], "local")
    generated = service.get(task["id"], "local")["exports"][0]
    assert generated["status"] == "completed", generated
    book = load_workbook(service.export_file(task["id"], run["id"], "local", "最终商品表.xlsx"))
    assert [cell.value for cell in book["平台商品"][1]] == ["商品名称", "货号", "销售价"]
    book.close()


def test_final_output_restores_grouped_headers_from_field_titles(tmp_path):
    from excel_auditor.workflow.models import OutputFile, OutputSheet

    config = default_config()
    fields = {field.rule.title: field.rule.name for field in config.fields}
    reference = tmp_path / "flat-output-template.xlsx"
    reference.write_bytes(workbook(rows=[], headers=["商品名称", "包装量 / 数值", "包装量 / 单位"], title="最终商品表"))
    spec = OutputFile(name="最终商品表.xlsx", sheets=[OutputSheet(
        name="最终商品表", header_row=1, data_start_row=2, columns=[
            OutputColumn(title="商品名称", field=fields["商品名称"]),
            OutputColumn(title="包装量 / 数值", field=fields["包装量/数值"]),
            OutputColumn(title="包装量 / 单位", field=fields["包装量/单位"]),
        ],
    )])
    destination = tmp_path / "grouped-output.xlsx"
    write_output(destination, spec, config, [{"values_json": {
        fields["商品名称"]: "现代简约客厅落地护眼灯", fields["包装量/数值"]: "2", fields["包装量/单位"]: "件",
    }}], reference)

    book = load_workbook(destination)
    try:
        sheet = book["最终商品表"]
        assert sheet["A1"].value == "商品名称"
        assert sheet["B1"].value == "包装量" and sheet["B2"].value == "数值" and sheet["C2"].value == "单位"
        assert {"A1:A2", "B1:C1"} <= {str(region) for region in sheet.merged_cells.ranges}
        assert sheet["A3"].value == "现代简约客厅落地护眼灯" and sheet["B3"].value == 2 and sheet["C3"].value == "件"
        assert sheet.freeze_panes == "A3"
        assert sheet.column_dimensions["A"].width >= 24
    finally:
        book.close()


@pytest.mark.parametrize("formula", ["=WEBSERVICE(\"https://example.org\")", "='[other.xlsx]Sheet1'!A1", "=SUM(A{bad})"])
def test_untrusted_formula_templates_rejected(formula):
    with pytest.raises(ValidationError): OutputColumn(title="计算", operation="formula", argument=formula)


def test_ai_secrets_remain_private_and_json_draft_is_validated(tmp_path):
    seen = []
    def transport(request):
        seen.append(request)
        if request.url.path.endswith("/models"):
            return httpx.Response(200, json={"data": [{"id": "gpt-5.6-luna"}]})
        payload = json.loads(request.content)
        assert payload["model"] == "gpt-5.6-luna"
        assert "samples" not in json.loads(payload["messages"][1]["content"])
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": '{"fields":[],"columns":[],"notes":["请确认业务规则"]}'}}]})
    client = AIClient(tmp_path, "local", httpx.MockTransport(transport))
    client.save(AISettings(base_url="https://model.example/v1", model="gpt-5.6-luna", api_key=SecretStr("test-private-secret")))
    assert "api_key" not in client.settings() and "secret" not in client.settings()
    assert client.models()["models"] == ["gpt-5.6-luna"]
    assert client.suggest(default_config(), "不要改价格", [{"product_name": "private sample"}])["notes"]
    client.save(AISettings(base_url="https://other.example/v1", model="gpt-5.6-luna"))
    assert not client.settings()["has_key"]


def test_concurrent_edits_have_exactly_one_winner(runtime):
    _, service, _ = runtime
    task = ready_task(runtime)
    row = service.rows(task["id"], "local", 0, 20)["items"][0]
    def write(value):
        try:
            service.edit(task["id"], "local", "local", EditRequest(base_revision=1, changes={row["id"]: {"stock": value}}))
            return "saved"
        except WorkflowError as exc:
            return exc.status
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(write, ["10", "20"]))
    assert sorted(map(str, results)) == ["409", "saved"]
    assert service.get(task["id"], "local")["revision"] == 2


def test_source_formulas_and_non_xlsx_are_not_silently_converted(runtime, tmp_path):
    task = create_task(runtime, workbook([["=1+1", "001", "2", "4", ""]]))
    assert task["status"] == "failed" and "公式" in task["payload"]["error"]
    path = tmp_path / "fake.xlsx"; path.write_bytes(b"not an excel file")
    with pytest.raises(WorkflowError): safe_book(path)


def test_api_full_sop_and_authorization(tmp_path, monkeypatch):
    monkeypatch.setenv("EXCEL_AUDITOR_DATA", str(tmp_path / "api"))
    for key in ("DATABASE_URL", "REDIS_URL", "S3_BUCKET", "EXCEL_AUDITOR_API_TOKEN", "EXCEL_AUDITOR_API_TOKENS_JSON"):
        monkeypatch.delenv(key, raising=False)
    import excel_auditor.api as module
    module = importlib.reload(module)
    with TestClient(module.app) as client:
        created_template = client.post("/api/v1/workflow-templates", json={"name": "API 流程"})
        assert created_template.status_code == 201, created_template.text
        draft = created_template.json()
        configured = client.put(f"/api/v1/workflow-templates/{draft['id']}", json={"base_revision": draft["revision"], "config": default_config().model_dump(mode="json"), "confirmed": True})
        assert configured.status_code == 200, configured.text
        published = client.post(f"/api/v1/workflow-templates/{draft['id']}/publish", json={"base_revision": configured.json()["revision"]})
        assert published.status_code == 200, published.text
        template = published.json()
        created = client.post("/api/v1/workflow-tasks", data={"template_id": template["id"]}, files={"excel_file": ("sample.xlsx", workbook())})
        assert created.status_code == 202, created.text
        task_id = created.json()["id"]
        task = client.get(f"/api/v1/workflow-tasks/{task_id}").json()
        ready = client.post(f"/api/v1/workflow-tasks/{task_id}/mapping", json=mapping(task).model_dump())
        assert ready.status_code == 200, ready.text
        downloaded = client.get(f"/api/v1/workflow-tasks/{task_id}/maintenance/download")
        assert downloaded.status_code == 200 and downloaded.content[:2] == b"PK"
        created_export = client.post(f"/api/v1/workflow-tasks/{task_id}/exports", json={"base_revision": 1})
        assert created_export.status_code == 202, created_export.text
        exported = client.get(f"/api/v1/workflow-tasks/{task_id}").json()
        assert exported["exports"][0]["status"] == "completed"
        run = exported["exports"][0]
        preview = client.get(f"/api/v1/workflow-tasks/{task_id}/exports/{run['id']}/preview/商品导入.xlsx")
        assert preview.status_code == 200
        preview_rows = preview.json()["sheets"][0]["rows"]
        header_depth = max(len(field["rule"]["title"].split("/")) for field in task["payload"]["config"]["fields"])
        assert preview_rows[header_depth][preview_rows[0].index("商品名称")] == "台灯"
        forbidden = client.post("/api/v1/workflow-templates", headers={"Origin": "https://evil.example"}, json={})
        assert forbidden.status_code == 403


def test_trial_applies_import_transforms_and_clears_previous_success(runtime):
    store, _, templates = runtime
    draft = templates.create("local")
    config = default_config()
    from excel_auditor.workflow.models import ImportRule
    config.import_rules = [ImportRule(field="sale_price", headers=["销售价"], operation="multiply", argument="2"),
                           ImportRule(field="product_name", headers=["商品名称", "货号"], operation="concat", argument="-")]
    draft = templates.update(draft["id"], "local", TemplateEdit(base_revision=0, config=config, confirmed=True))
    trial = templates.trial(draft["id"], "local", draft["revision"], workbook())
    book = load_workbook(store.directory(draft["id"]) / "trial-0.xlsx")
    values = list(book.active.values)
    book.close()
    header_depth = max(len(field.rule.title.split("/")) for field in config.fields)
    assert values[header_depth][values[0].index("商品名称")] == "台灯-0001"
    price_index = next(i for i, column in enumerate(config.outputs[0].sheets[0].columns) if column.field == "sale_price")
    assert values[header_depth][price_index] == 59
    with pytest.raises(WorkflowError):
        templates.trial(draft["id"], "local", trial["revision"], b"broken")
    failed = templates.get(draft["id"], "local")
    assert not failed["payload"].get("tested_hash")
    published = templates.publish(draft["id"], "local", failed["revision"])
    assert published["status"] == "published"


def test_reference_header_changes_grouping_and_sku_expansion(tmp_path):
    from excel_auditor.workflow.models import OutputFile, OutputSheet
    config = default_config()
    reference = tmp_path / "reference.xlsx"
    reference.write_bytes(workbook(title="输出"))
    spec = OutputFile(sheets=[OutputSheet(name="输出", group_by="brand", sku_fields=["sku"], columns=[
        OutputColumn(title="编码", field="sku"), OutputColumn(title="新名称", field="product_name"),
        OutputColumn(title="公式", operation="formula", argument="=LEN(A{row})")])])
    destination = tmp_path / "result.xlsx"
    result = write_output(destination, spec, config, [{"values_json": {"product_name": "台灯", "brand": "测试", "sku": "001|002"}}], reference)
    assert result["sheets"][0]["records"] == 2
    book = load_workbook(destination)
    assert book.sheetnames == ["输出-测试"]
    assert [c.value for c in book.active[1]][:3] == ["编码", "新名称", "公式"]
    assert book.active["A2"].value == "001" and book.active["A3"].value == "002"
    assert book.active["C3"].value == "=LEN(A3)"
    assert not any(cell.value for row in book.active.iter_rows(min_col=4) for cell in row)
    book.close()


def test_numeric_sort_and_readonly_fields(runtime):
    _, service, _ = runtime
    task = ready_task(runtime, content=workbook([["A", "1", "100", "2", ""], ["B", "2", "20", "10", ""]]))
    rows = service.rows(task["id"], "local", 0, 30, sort="sale_price")["items"]
    assert [row["values_json"]["sale_price"] for row in rows] == ["20", "100"]
    with pytest.raises(WorkflowError, match="过长"):
        service.edit(task["id"], "local", "local", EditRequest(base_revision=1, additions=[{"product_name": "x" * 32768}]))


def test_template_fork_preserves_purpose(runtime):
    _, _, templates = runtime
    original = templates.upload("local", "维护模板", workbook(), "maintenance")
    templates.parse(original["id"], "local")
    forked = templates.create("local", original["id"])
    assert forked["payload"]["purpose"] == "maintenance"
